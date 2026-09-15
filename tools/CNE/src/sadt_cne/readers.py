"""Finding the notes, and turning each of them into plain text.

Ported from `CNE_CLI/CNE_CLI.py`'s `clean_text` / `extract_text*` and its
discovery glob. What differs, and why:

* **discovery is recursive and case-insensitive.** Upstream globbed
  `<folder>/*.txt`, `*.pdf`, `*.docx` (`:132-133`), so `NOTE.PDF` and every note
  filed in a per-patient subfolder were invisible -- not refused, not counted,
  not mentioned. A cohort is filed in subfolders; that is what a cohort is.
* **a `.txt` is decoded, not assumed.** Upstream opened it UTF-8 strict with no
  `errors=` (`:80`), so a note exported from Word on Windows raised
  `UnicodeDecodeError` inside the per-file `except`, was counted as a failure
  nobody saw, and the run still reported success. The encoding that worked is
  recorded per note.
* **a `.docx` is read whole.** Upstream read `doc.paragraphs` only, which drops
  every table, header and footer. Clinical notes are frequently tabular, and a
  dropped table is a silently shorter note the model then extracts from.
* **an empty extraction is not text.** A scanned, image-only PDF has no text
  layer and yields `""`. Upstream sent that to the model, which answered with a
  plausible extraction that was written to disk and counted as a success. The
  guard is in `__init__.run`, and this module's job is to report honestly what
  it got.
"""

import os
import unicodedata
from pathlib import Path

from .dependencies import require

# The three formats this tool advertises, which are the three it reads. Matched
# case-insensitively: `.PDF` off a Windows share is the same note.
NOTE_EXTENSIONS = (".txt", ".pdf", ".docx")

# Tried in order for a `.txt`, and the first that decodes cleanly wins.
# `utf-8-sig` reads UTF-8 with or without a byte-order mark and strips it, so a
# BOM never becomes a stray character at the head of the note; `cp1252` is what
# Word and Excel write on Windows; `latin-1` decodes any byte sequence at all
# and is the floor, so this never raises. Which one was used is reported,
# because a note that only decoded as latin-1 is one to look at.
TEXT_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")

# Characters that vary between PDF and Word exports but never appeared in the
# plain-text notes the models were fine-tuned on. Normalising them keeps the
# text the model receives identical whatever the source format. Ported verbatim
# from upstream's `_TEXT_REPLACEMENTS`.
_TEXT_REPLACEMENTS = {
    "’": "'",   # right single quotation mark
    "‘": "'",   # left single quotation mark
    "“": '"',   # left double quotation mark
    "”": '"',   # right double quotation mark
    "–": "-",   # en dash
    "—": "-",   # em dash
    " ": " ",   # non-breaking space
}

# Names never treated as a note: dotfiles, macOS archive residue, and the
# `~$name.docx` lock file Word leaves beside an open document -- which has the
# `.docx` extension, is not a document, and makes python-docx raise.
_SKIPPED_PREFIXES = ("~$", ".")
_SKIPPED_DIRECTORIES = ("__MACOSX",)


def normalise(text: str) -> str:
    """One note's text, in the shape the model was trained to read."""
    for old, new in _TEXT_REPLACEMENTS.items():
        text = text.replace(old, new)
    # Collapse the runs of blank lines that PDF and Word extraction introduce,
    # keeping single ones: a paragraph break carries meaning, twelve do not.
    cleaned = []
    for line in text.splitlines():
        line = line.strip()
        if line or (cleaned and cleaned[-1]):
            cleaned.append(line)
    return "\n".join(cleaned).strip()


def extension_of(name: str) -> str:
    """The note extension of `name`, lowercased, or `""`."""
    suffix = os.path.splitext(name)[1].lower()
    return suffix if suffix in NOTE_EXTENSIONS else ""


def is_note_file(name: str) -> bool:
    """Whether a file NAME is one this tool reads. Case-insensitive."""
    base = os.path.basename(name)
    if base.startswith(_SKIPPED_PREFIXES):
        return False
    return bool(extension_of(base))


def discover_notes(root, excluded=()) -> list:
    """Every note under `root`, sorted, as a list of `Path`.

    `root` may be one note file or a folder; a folder is walked recursively,
    which is the fix for upstream's single-level glob. `excluded` names
    directories whose contents are not notes -- it is how the output folder is
    kept out of the input when a caller nests one inside the other, so a second
    run cannot extract from the first run's extractions.

    Returns files only. Sorted by relative POSIX path so two runs on the same
    folder process the notes in the same order, which is what makes a report
    comparable between runs.
    """
    root = Path(root)
    if root.is_file():
        return [root] if is_note_file(root.name) else []
    if not root.is_dir():
        return []

    excluded = [Path(path).resolve() for path in excluded]
    found = []
    for path in root.rglob("*"):
        if not path.is_file() or not is_note_file(path.name):
            continue
        parts = path.relative_to(root).parts[:-1]
        if any(part.startswith(_SKIPPED_PREFIXES) or part in _SKIPPED_DIRECTORIES
               for part in parts):
            continue
        if any(_is_within(path, directory) for directory in excluded):
            continue
        found.append(path)
    return sorted(found, key=lambda path: path.relative_to(root).as_posix())


def _is_within(path: Path, directory: Path) -> bool:
    """Whether `path` sits inside `directory` (or is it)."""
    try:
        path.resolve().relative_to(directory)
    except ValueError:
        return False
    return True


def read_note(path) -> tuple:
    """One note as `(text, details)`, whatever its format.

    `details` says how it was read -- the reader, and for a `.txt` the encoding
    that decoded it -- and goes into the report. Reading is deliberately
    separate from deciding what to do with an empty result: this returns `""`
    for a PDF with no text layer rather than raising, and `run` is what refuses
    to send it to the model.
    """
    path = Path(path)
    extension = extension_of(path.name)
    if extension == ".pdf":
        return read_pdf(path), {"reader": "pdf"}
    if extension == ".docx":
        return read_docx(path), {"reader": "docx"}
    if extension == ".txt":
        text, encoding = read_text(path)
        return text, {"reader": "txt", "encoding": encoding}
    raise ValueError(
        f"'{path.name}' is not a note CNE reads. Supported: "
        f"{', '.join(NOTE_EXTENSIONS)}."
    )


def read_text(path) -> tuple:
    """A plain-text note as `(text, encoding)`, trying TEXT_ENCODINGS in order.

    Never raises `UnicodeDecodeError`: `latin-1` decodes every byte sequence,
    so the worst case is mojibake reported as such rather than a note counted
    as a phantom success.
    """
    raw = Path(path).read_bytes()
    for encoding in TEXT_ENCODINGS:
        try:
            return normalise(raw.decode(encoding)), encoding
        except UnicodeDecodeError:
            continue
    # Unreachable while latin-1 is last -- kept so removing it from the tuple
    # fails loudly here instead of raising from inside the loop.
    raise ValueError(f"'{Path(path).name}' could not be decoded as text.")


def read_pdf(path) -> str:
    """A PDF's text layer. An image-only scan legitimately has none."""
    fitz = require("fitz", "pymupdf")
    document = fitz.open(str(path))
    try:
        text = "\n".join(page.get_text() for page in document)
    finally:
        document.close()
    return normalise(text)


def read_docx(path) -> str:
    """A Word document: body paragraphs, tables, headers and footers.

    Upstream read `doc.paragraphs`, which is the body's paragraphs and nothing
    else -- every table cell, every running header and every footer was dropped
    without a word. The body is walked in document order rather than
    paragraphs-then-tables, so a table stays attached to the sentence that
    introduces it.
    """
    docx = require("docx", "python-docx")
    document = docx.Document(str(path))

    pieces = []
    for section in document.sections:
        for part in (section.header, section.footer):
            pieces.extend(_docx_container_text(part))
    pieces.extend(_docx_body_text(document))
    return normalise("\n".join(pieces))


def _docx_body_text(document) -> list:
    """The document body, paragraphs and tables interleaved in reading order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    pieces = []
    for child in body.iterchildren():
        tag = _local_name(child.tag)
        if tag == "p":
            pieces.append(Paragraph(child, document).text)
        elif tag == "tbl":
            pieces.extend(_docx_table_text(Table(child, document)))
    return pieces


def _docx_container_text(container) -> list:
    """A header or footer: its paragraphs, then its tables."""
    pieces = [paragraph.text for paragraph in container.paragraphs]
    for table in container.tables:
        pieces.extend(_docx_table_text(table))
    return pieces


def _docx_table_text(table) -> list:
    """One table, a row per line, cells separated by a tab.

    A tab rather than a newline because a row of a clinical table is one
    record; splitting it across lines is what makes "Pain 7" read as two
    unrelated notes.
    """
    lines = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        line = "\t".join(cell for cell in cells if cell)
        if line:
            lines.append(line)
    return lines


def _local_name(tag) -> str:
    """`p` from `{...wordprocessingml...}p`."""
    text = str(tag)
    return text.rsplit("}", 1)[-1]


def looks_empty(text: str) -> bool:
    """Whether there is nothing here for the model to read.

    Whitespace-only is empty, and so is a page of control characters: a text
    layer that carries no letter or digit is not a note, and sending it to the
    model is how upstream produced a plausible extraction from a scanned page.
    """
    return not any(
        character.isalnum() or unicodedata.category(character).startswith("L")
        for character in text
    )
