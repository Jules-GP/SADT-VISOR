"""CNE, with the language model stubbed.

No weights, no GPU, no network. Everything AROUND the model runs for real:
discovery, the three readers, encodings, output naming, the per-note failure
path, the report and the tally. The model is the one thing replaced, because a
4.4 GB GGUF is not a test fixture.

Every test is named after the behaviour it pins; the ones that pin a defect
found in `CNE_CLI/CNE_CLI.py` or `CNE/CNE.py` say so in their docstring.
"""

import ctypes
import json
import os
import subprocess
import sys
import types
from pathlib import Path

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))

import sadt_cne
from sadt_cne import dependencies, extraction, readers
from sadt_cne.dependencies import ToolUnavailableError

ANSWER = '{"chief_complaint": "jaw pain", "side": "left"}'


# --------------------------------------------------------------------------
# Fixtures: notes on disk in each of the three formats, and a stubbed engine.
# --------------------------------------------------------------------------

def write_txt(path, text="Patient reports left jaw pain.", encoding="utf-8"):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode(encoding))
    return path


def write_pdf(path, text="Patient reports left jaw pain."):
    """A real PDF with a real text layer, written by the reader's own library."""
    import fitz

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    document.save(str(path))
    document.close()
    return path


def write_docx(path, paragraphs=("Patient reports left jaw pain.",),
               table=None, header=None, footer=None):
    """A real .docx, written by python-docx."""
    import docx

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        built = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, cell in enumerate(row):
                built.cell(row_index, column_index).text = cell
    if header:
        document.sections[0].header.paragraphs[0].text = header
    if footer:
        document.sections[0].footer.paragraphs[0].text = footer
    document.save(str(path))
    return path


@pytest.fixture
def model(tmp_path):
    """A hosted bundle folder holding one .gguf, as DATA/CNE/models/TMJ/ is."""
    bundle = tmp_path / "models" / "TMJ"
    bundle.mkdir(parents=True)
    (bundle / "qwen-ft-q4_k_m.gguf").write_bytes(b"GGUF not really")
    return bundle


@pytest.fixture
def llm(monkeypatch):
    """Replace the engine. Everything else in the run is the real thing.

    The returned recorder is both the stand-in engine and the log: `loaded`
    holds the arguments the loader was given, `calls` every completion asked
    for, and `answers`/`finish_reason` let a test choose what comes back.
    """
    recorder = types.SimpleNamespace(
        loaded=None, calls=[], answers=ANSWER, finish_reason="stop",
        gpu_offload=False,
    )

    def load_model(model_file, context_tokens, seed, n_gpu_layers):
        recorder.loaded = {
            "model_file": Path(model_file),
            "context_tokens": context_tokens,
            "seed": seed,
            "n_gpu_layers": n_gpu_layers,
        }
        return recorder

    def complete(engine, messages, max_tokens, temperature):
        recorder.calls.append({
            "messages": messages,
            "max_tokens": max_tokens,
            "temperature": temperature,
        })
        answers = recorder.answers
        if callable(answers):
            return answers(messages)
        if isinstance(answers, list):
            answers = answers[len(recorder.calls) - 1]
        if isinstance(answers, tuple):
            return answers
        return answers, recorder.finish_reason

    monkeypatch.setattr(sadt_cne, "load_model", load_model)
    monkeypatch.setattr(sadt_cne, "complete", complete)
    # Stubbed like the rest of the engine: whether the wheel installed in THIS
    # checkout has CUDA compiled in must not decide what the suite asserts.
    monkeypatch.setattr(
        sadt_cne, "supports_gpu_offload", lambda: recorder.gpu_offload
    )
    return recorder


def report_of(output_dir):
    return json.loads((Path(output_dir) / "CNE_report.json").read_text())


def entry_for(output_dir, name):
    for entry in report_of(output_dir)["notes"]:
        if entry["input"] == name:
            return entry
    raise AssertionError(f"no report entry for {name}")


# --------------------------------------------------------------------------
# Discovery (defect: `glob.glob(folder + "*.ext")`, CNE_CLI.py:132-133)
# --------------------------------------------------------------------------

def test_notes_are_discovered_recursively(tmp_path):
    """Upstream globbed one level, so every note filed per patient was
    invisible -- not refused, not counted, not mentioned."""
    write_txt(tmp_path / "top.txt")
    write_txt(tmp_path / "patient_1" / "visit_a" / "deep.txt")

    found = readers.discover_notes(tmp_path)

    assert [path.relative_to(tmp_path).as_posix() for path in found] == [
        "patient_1/visit_a/deep.txt", "top.txt",
    ]


def test_discovery_is_case_insensitive_about_extensions(tmp_path):
    """`NOTE.PDF` off a Windows share is the same note."""
    write_txt(tmp_path / "SHOUTED.TXT")
    write_txt(tmp_path / "mixed.TxT")

    found = readers.discover_notes(tmp_path)

    assert sorted(path.name for path in found) == ["SHOUTED.TXT", "mixed.TxT"]


def test_discovery_returns_files_never_directories(tmp_path):
    (tmp_path / "folder.txt").mkdir()
    write_txt(tmp_path / "real.txt")

    found = readers.discover_notes(tmp_path)

    assert [path.name for path in found] == ["real.txt"]
    assert all(path.is_file() for path in found)


def test_discovery_is_sorted_so_two_runs_agree(tmp_path):
    for name in ("z.txt", "a.txt", "m/b.txt", "m/a.txt"):
        write_txt(tmp_path / name)

    found = readers.discover_notes(tmp_path)

    assert [path.relative_to(tmp_path).as_posix() for path in found] == [
        "a.txt", "m/a.txt", "m/b.txt", "z.txt",
    ]


@pytest.mark.parametrize("name,expected", [
    ("note.txt", True), ("note.pdf", True), ("note.docx", True),
    ("NOTE.PDF", True), ("note.DOCX", True),
    ("note.doc", False), ("note.rtf", False), ("note.odt", False),
    ("note", False), ("note.txt.bak", False),
])
def test_only_the_three_advertised_formats_are_notes(name, expected):
    assert readers.is_note_file(name) is expected


def test_dotfiles_are_not_notes(tmp_path):
    write_txt(tmp_path / ".hidden.txt")
    write_txt(tmp_path / "visible.txt")

    found = readers.discover_notes(tmp_path)

    assert [path.name for path in found] == ["visible.txt"]


def test_word_lock_files_are_not_notes(tmp_path):
    """`~$note.docx` has the extension, is not a document, and makes
    python-docx raise -- so it would be reported as a failed patient."""
    write_txt(tmp_path / "~$note.docx")
    write_txt(tmp_path / "note.txt")

    found = readers.discover_notes(tmp_path)

    assert [path.name for path in found] == ["note.txt"]


def test_archive_residue_directories_are_skipped(tmp_path):
    write_txt(tmp_path / "__MACOSX" / "ghost.txt")
    write_txt(tmp_path / "note.txt")

    found = readers.discover_notes(tmp_path)

    assert [path.name for path in found] == ["note.txt"]


def test_a_single_note_file_is_accepted_as_input(tmp_path, model, llm):
    note = write_txt(tmp_path / "one.txt")

    output = sadt_cne.run(note, "TMJ", model, tmp_path / "out")

    assert (output / "Extraction_one.txt.json").is_file()


def test_an_empty_folder_is_refused_not_reported_as_a_finished_run(tmp_path, model, llm):
    """Upstream exited 0 here (CNE_CLI.py:135-139), so a mistyped path reached
    the user as a completed extraction."""
    (tmp_path / "notes").mkdir()

    with pytest.raises(ValueError, match="No note found"):
        sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")


def test_a_folder_of_unsupported_files_is_refused(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "scan.dcm")

    with pytest.raises(ValueError, match="No note found"):
        sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")


def test_a_missing_input_path_is_a_file_not_found_error(tmp_path, model, llm):
    with pytest.raises(FileNotFoundError):
        sadt_cne.run(tmp_path / "nowhere", "TMJ", model, tmp_path / "out")


# --------------------------------------------------------------------------
# Output naming (defect: Extraction_<stem>.txt, CNE_CLI.py:226-233)
# --------------------------------------------------------------------------

def test_three_notes_sharing_a_stem_write_three_extractions(tmp_path, model, llm):
    """The destroying defect. `B_001.txt`, `B_001.pdf` and `B_001.docx` in one
    folder all wrote `Extraction_B_001.txt` with `open(..., "w")`: two of the
    three extractions were lost and the run reported 3/3 processed."""
    notes = tmp_path / "notes"
    write_txt(notes / "B_001.txt", "Text note, left side.")
    write_pdf(notes / "B_001.pdf", "PDF note, right side.")
    write_docx(notes / "B_001.docx", ["Word note, both sides."])

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    written = sorted(path.name for path in output.glob("Extraction_*.json"))
    assert written == [
        "Extraction_B_001.docx.json",
        "Extraction_B_001.pdf.json",
        "Extraction_B_001.txt.json",
    ]
    assert report_of(output)["summary"] == "3/3 note(s) extracted"


def test_the_output_name_keeps_the_source_extension(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "B_001.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert (output / "Extraction_B_001.txt.txt").is_file()
    assert (output / "Extraction_B_001.txt.json").is_file()


def test_subfolders_are_mirrored_in_the_output(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "patient_1" / "visit_a" / "note.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert (output / "patient_1" / "visit_a" / "Extraction_note.txt.json").is_file()


def test_two_notes_of_one_name_in_different_folders_do_not_collide(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "patient_1" / "note.txt", "First patient.")
    write_txt(tmp_path / "notes" / "patient_2" / "note.txt", "Second patient.")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert (output / "patient_1" / "Extraction_note.txt.json").is_file()
    assert (output / "patient_2" / "Extraction_note.txt.json").is_file()
    assert report_of(output)["summary"] == "2/2 note(s) extracted"


def test_planned_output_paths_are_checked_for_collisions(tmp_path):
    """Belt and braces on the naming rule: the failure it guards destroys one
    patient's extraction while reporting success."""
    root = tmp_path / "notes"
    write_txt(root / "a.txt")

    planned = sadt_cne.plan_outputs(
        [root / "a.txt"], root, tmp_path / "out"
    )

    assert len(planned) == 1
    with pytest.raises(ValueError, match="would both write"):
        sadt_cne.plan_outputs(
            [root / "a.txt", root / "a.txt"], root, tmp_path / "out"
        )


def test_a_json_is_written_beside_every_text_extraction(tmp_path, model, llm):
    """Upstream flattened the model's JSON to `key : value` text at the last
    step and threw the object away, so nothing downstream could read it."""
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert json.loads((output / "Extraction_a.txt.json").read_text()) == {
        "chief_complaint": "jaw pain", "side": "left",
    }


def test_the_text_extraction_is_key_colon_value(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert (output / "Extraction_a.txt.txt").read_text() == (
        "chief_complaint : jaw pain\nside : left\n"
    )


# --------------------------------------------------------------------------
# The output directory (defect: created by the panel, CNE/CNE.py:578)
# --------------------------------------------------------------------------

def test_the_engine_creates_its_own_output_directory(tmp_path, model, llm):
    """Upstream's was created by the Slicer panel and never by the engine, so
    headless every write raised FileNotFoundError, was swallowed per file, and
    the run still claimed success."""
    write_txt(tmp_path / "notes" / "a.txt")
    destination = tmp_path / "does" / "not" / "exist"

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, destination)

    assert output.is_dir()
    assert (output / "Extraction_a.txt.json").is_file()


def test_mirrored_subdirectories_are_created_too(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "deep" / "deeper" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert (output / "deep" / "deeper" / "Extraction_a.txt.json").is_file()


# --------------------------------------------------------------------------
# Input and output folders (defect: a run reading its own extractions)
# --------------------------------------------------------------------------

def test_writing_into_the_notes_folder_is_refused(tmp_path, model, llm):
    """Upstream allowed it, and its outputs matched its own discovery glob."""
    notes = tmp_path / "notes"
    write_txt(notes / "a.txt")

    with pytest.raises(ValueError, match="must not be the notes folder"):
        sadt_cne.run(notes, "TMJ", model, notes)


def test_an_output_folder_inside_the_notes_folder_is_not_discovered(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_txt(notes / "a.txt")
    write_txt(notes / "extractions" / "leftover.txt", "A previous run's output.")

    output = sadt_cne.run(notes, "TMJ", model, notes / "extractions")

    assert [entry["input"] for entry in report_of(output)["notes"]] == ["a.txt"]


def test_a_second_run_extracts_the_same_notes_not_its_own_output(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_txt(notes / "a.txt")
    output_dir = notes / "extractions"

    sadt_cne.run(notes, "TMJ", model, output_dir)
    second = sadt_cne.run(notes, "TMJ", model, output_dir)

    assert report_of(second)["summary"] == "1/1 note(s) extracted"


# --------------------------------------------------------------------------
# No text (the highest-severity defect: a scanned PDF hallucinated into a result)
# --------------------------------------------------------------------------

def test_a_note_with_no_text_is_never_sent_to_the_model(tmp_path, model, llm):
    """THE defect. An image-only PDF yields "", which upstream sent to the model
    anyway; the model answered with a plausible extraction from nothing, and it
    was written to disk and counted as a success."""
    notes = tmp_path / "notes"
    write_pdf(notes / "scanned.pdf", text="")
    write_txt(notes / "real.txt")

    sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    sent = [call["messages"][-1]["content"] for call in llm.calls]
    assert sent == ["Patient reports left jaw pain."]


def test_a_note_with_no_text_writes_nothing(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_pdf(notes / "scanned.pdf", text="")
    write_txt(notes / "real.txt")

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert not (output / "Extraction_scanned.pdf.json").exists()
    assert not (output / "Extraction_scanned.pdf.txt").exists()


def test_a_note_with_no_text_is_reported_as_failed_and_says_why(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_pdf(notes / "scanned.pdf", text="")
    write_txt(notes / "real.txt")

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    entry = entry_for(output, "scanned.pdf")
    assert entry["status"] == "failed"
    assert "no text could be extracted" in entry["reason"]
    assert "OCR" in entry["reason"]


def test_a_whitespace_only_note_counts_as_empty(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_txt(notes / "blank.txt", "\n\n   \t\n")
    write_txt(notes / "real.txt")

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert entry_for(output, "blank.txt")["status"] == "failed"


@pytest.mark.parametrize("text,empty", [
    ("", True), ("   \n\t ", True), ("...---...", True),
    ("Pain", False), ("7", False), ("Douleur à gauche", False),
])
def test_a_note_carrying_no_letter_or_digit_is_empty(text, empty):
    assert readers.looks_empty(text) is empty


def test_the_character_count_of_every_note_is_reported(tmp_path, model, llm):
    """A near-empty note is legitimate and is also the signature of a scan that
    barely OCR'd; only the caller can tell them apart, so the number is
    published rather than judged."""
    write_txt(tmp_path / "notes" / "a.txt", "Pain 7/10.")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert entry_for(output, "a.txt")["characters"] == len("Pain 7/10.")


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------

def test_a_plain_text_note_is_read(tmp_path):
    path = write_txt(tmp_path / "a.txt", "Pain on the left.")

    text, details = readers.read_note(path)

    assert text == "Pain on the left."
    assert details == {"reader": "txt", "encoding": "utf-8-sig"}


def test_a_windows_1252_note_is_read_instead_of_becoming_a_phantom_success(tmp_path):
    """Upstream opened `.txt` UTF-8 strict with no `errors=` (CNE_CLI.py:80), so
    a note exported from Word raised inside the per-file `except`, was counted
    as a failure nobody saw, and the run still reported success."""
    path = write_txt(tmp_path / "a.txt", "Café — naïve", encoding="cp1252")

    with pytest.raises(UnicodeDecodeError):
        path.read_bytes().decode("utf-8")

    text, details = readers.read_note(path)

    assert "Café" in text and "naïve" in text
    assert details["encoding"] == "cp1252"


def test_a_utf8_bom_is_stripped_rather_than_read_as_a_character(tmp_path):
    path = write_txt(tmp_path / "a.txt", "﻿Pain on the left.")

    text, _details = readers.read_note(path)

    assert text == "Pain on the left."


def test_the_encoding_that_worked_is_reported(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt", "Café", encoding="cp1252")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert entry_for(output, "a.txt")["encoding"] == "cp1252"


def test_undecodable_bytes_never_raise(tmp_path):
    """latin-1 is the floor, so the worst case is mojibake reported as such."""
    path = tmp_path / "a.txt"
    path.write_bytes(bytes(range(0x80, 0x100)) + b"Pain")

    text, details = readers.read_note(path)

    assert "Pain" in text
    assert details["encoding"] in readers.TEXT_ENCODINGS


def test_typographic_characters_are_normalised(tmp_path):
    path = write_txt(tmp_path / "a.txt", "the patient’s “pain” – 7/10")

    text, _details = readers.read_note(path)

    assert text == "the patient's \"pain\" - 7/10"


def test_runs_of_blank_lines_are_collapsed_to_one(tmp_path):
    path = write_txt(tmp_path / "a.txt", "First\n\n\n\n\nSecond")

    text, _details = readers.read_note(path)

    assert text == "First\n\nSecond"


def test_a_pdf_text_layer_is_read(tmp_path):
    path = write_pdf(tmp_path / "a.pdf", "Pain on the left.")

    text, details = readers.read_note(path)

    assert "Pain on the left." in text
    assert details == {"reader": "pdf"}


def test_a_pdf_with_no_text_layer_reads_as_empty_rather_than_raising(tmp_path):
    """This is what a scanned page is: a picture, and no text at all."""
    path = write_pdf(tmp_path / "a.pdf", text="")

    text, _details = readers.read_note(path)

    assert text == ""


def test_docx_paragraphs_are_read(tmp_path):
    path = write_docx(tmp_path / "a.docx", ["First line.", "Second line."])

    text, details = readers.read_note(path)

    assert "First line." in text and "Second line." in text
    assert details == {"reader": "docx"}


def test_docx_tables_are_read(tmp_path):
    """Upstream read `doc.paragraphs` only. Clinical notes are frequently
    tabular, and a dropped table is a silently shorter note."""
    path = write_docx(
        tmp_path / "a.docx",
        paragraphs=["Findings:"],
        table=[["Site", "Score"], ["Left TMJ", "7"]],
    )

    text, _details = readers.read_note(path)

    assert "Left TMJ" in text and "7" in text


def test_a_docx_table_row_stays_on_one_line(tmp_path):
    """A row is one record. Split across lines, "Pain 7" reads as two notes."""
    path = write_docx(
        tmp_path / "a.docx", paragraphs=[""],
        table=[["Left TMJ", "7"]],
    )

    text, _details = readers.read_note(path)

    assert "Left TMJ\t7" in text


def test_docx_headers_and_footers_are_read(tmp_path):
    path = write_docx(
        tmp_path / "a.docx", paragraphs=["Body text."],
        header="Patient B_001", footer="Page 1 of 1",
    )

    text, _details = readers.read_note(path)

    assert "Patient B_001" in text
    assert "Page 1 of 1" in text


def test_docx_body_content_is_read_in_document_order(tmp_path):
    """Upstream would have read paragraphs then tables, detaching a table from
    the sentence that introduces it."""
    import docx

    path = tmp_path / "a.docx"
    document = docx.Document()
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell"
    document.add_paragraph("After the table.")
    document.save(str(path))

    text, _details = readers.read_note(path)

    assert text.index("Before the table.") < text.index("Cell")
    assert text.index("Cell") < text.index("After the table.")


def test_reading_a_format_cne_does_not_know_is_a_value_error(tmp_path):
    path = write_txt(tmp_path / "a.rtf")

    with pytest.raises(ValueError, match="not a note CNE reads"):
        readers.read_note(path)


def test_an_unreadable_note_fails_only_that_note(tmp_path, model, llm, monkeypatch):
    """Upstream wrapped the whole loop in `except Exception: sys.exit(1)`
    (CNE_CLI.py:248-251), so the first bad file ended the cohort."""
    notes = tmp_path / "notes"
    write_txt(notes / "bad.txt")
    write_txt(notes / "good.txt")

    real = readers.read_note

    def flaky(path):
        if Path(path).name == "bad.txt":
            raise OSError("disk went away")
        return real(path)

    monkeypatch.setattr(sadt_cne, "read_note", flaky)
    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert report_of(output)["summary"] == "1/2 note(s) extracted"
    assert entry_for(output, "good.txt")["status"] == "ok"
    assert "disk went away" in entry_for(output, "bad.txt")["reason"]


# --------------------------------------------------------------------------
# Dependencies (defect: sys.exit(1) for a reader an all-.txt batch never uses)
# --------------------------------------------------------------------------

def test_a_text_only_batch_needs_neither_pdf_nor_word_reader(tmp_path, model, llm, monkeypatch):
    """Upstream exited 1 when `pymupdf` OR `python-docx` was missing
    (CNE_CLI.py:105-115), whatever the batch actually held."""
    def refuse(name, package=""):
        raise AssertionError(f"a .txt batch must not import {name}")

    monkeypatch.setattr(readers, "require", refuse)
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert report_of(output)["summary"] == "1/1 note(s) extracted"


def test_a_missing_pdf_reader_costs_the_pdf_notes_and_nothing_else(tmp_path, model, llm, monkeypatch):
    def no_fitz(name):
        raise ImportError(f"No module named {name!r}")

    monkeypatch.setattr(dependencies.importlib, "import_module", no_fitz)
    notes = tmp_path / "notes"
    write_pdf(notes / "a.pdf")
    write_txt(notes / "b.txt")

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert entry_for(output, "b.txt")["status"] == "ok"
    assert "pymupdf" in entry_for(output, "a.pdf")["reason"]


def test_a_pdf_only_batch_without_a_pdf_reader_is_unavailable_not_invalid(tmp_path, model, llm, monkeypatch):
    """A deployment problem answers 503, not 422: nothing the caller sends
    would make a missing package appear."""
    def no_fitz(name):
        raise ImportError(f"No module named {name!r}")

    monkeypatch.setattr(dependencies.importlib, "import_module", no_fitz)
    write_pdf(tmp_path / "notes" / "a.pdf")

    with pytest.raises(ToolUnavailableError, match="pymupdf"):
        sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")


def test_a_missing_llama_cpp_is_a_tool_unavailable_error(tmp_path, monkeypatch):
    def no_llama(name):
        raise ImportError(f"No module named {name!r}")

    monkeypatch.setattr(dependencies.importlib, "import_module", no_llama)

    with pytest.raises(ToolUnavailableError, match="llama-cpp-python"):
        extraction.load_model(tmp_path / "m.gguf", 2048, 0, 0)


def test_require_returns_the_module_when_it_is_there():
    assert dependencies.require("json", "json") is json


# --------------------------------------------------------------------------
# The note type (defect: `== "TMJ"` at :150, `.upper() == "TMJ"` at :193)
# --------------------------------------------------------------------------

def test_the_tmj_prompt_and_context_are_chosen_from_one_table(tmp_path, model, llm):
    """Upstream compared the type case-sensitively for the context and
    case-insensitively for the prompt, so `"tmj"` got the TMJ instruction with
    the Ortho 2048 window and long notes were silently truncated."""
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert llm.loaded["context_tokens"] == 6144
    assert llm.calls[0]["messages"][0] == {
        "role": "system", "content": extraction.INSTRUCTION_TMJ,
    }


def test_ortho_gets_no_system_prompt_and_its_own_context(tmp_path, model, llm):
    """Its model is fine-tuned to answer from the note alone; adding an
    instruction it never saw in training changes what it emits."""
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "Ortho", model, tmp_path / "out")

    assert llm.loaded["context_tokens"] == 2048
    assert [message["role"] for message in llm.calls[0]["messages"]] == ["user"]


def test_the_note_type_table_and_the_prompt_table_cover_the_same_types():
    assert set(extraction.CONTEXT_TOKENS) == set(extraction.SYSTEM_PROMPTS)


def test_an_explicit_context_size_overrides_the_note_type_default(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                 context_tokens=12288)

    assert llm.loaded["context_tokens"] == 12288


def test_a_negative_context_size_is_refused():
    with pytest.raises(ValueError, match="context_tokens"):
        extraction.context_for("TMJ", -1)


def test_the_note_type_is_recorded_in_the_report(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "Ortho", model, tmp_path / "out")

    assert report_of(output)["notes_type"] == "Ortho"
    assert report_of(output)["system_prompt"] is False


# --------------------------------------------------------------------------
# Generation (defect: max_tokens=500 truncating, temperature=0.1 drifting)
# --------------------------------------------------------------------------

def test_a_truncated_answer_is_a_failure_not_a_written_result(tmp_path, model, llm):
    """Upstream capped generation at 500 tokens; the JSON that ran off the end
    failed to parse and the fragment was written as if it were the result."""
    write_txt(tmp_path / "notes" / "a.txt")
    write_txt(tmp_path / "notes" / "b.txt")
    llm.answers = [('{"pain": "yes"', "length"), (ANSWER, "stop")]

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert entry_for(output, "a.txt")["status"] == "failed"
    assert entry_for(output, "b.txt")["status"] == "ok"
    assert "cut off at max_tokens" in entry_for(output, "a.txt")["reason"]


def test_a_truncated_answer_writes_nothing(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")
    write_txt(tmp_path / "notes" / "b.txt")
    llm.answers = [('{"pain": "yes"', "length"), (ANSWER, "stop")]

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert not (output / "Extraction_a.txt.txt").exists()
    assert not (output / "Extraction_a.txt.json").exists()


def test_the_default_temperature_is_zero_so_a_run_can_be_checked(tmp_path, model, llm):
    """Upstream used 0.1, so the same note extracted twice gave two answers."""
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert llm.calls[0]["temperature"] == 0.0


def test_max_tokens_reaches_the_model(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                 max_tokens=64)

    assert llm.calls[0]["max_tokens"] == 64


def test_the_seed_reaches_the_engine_and_the_report(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                          seed=17, temperature=0.7)

    assert llm.loaded["seed"] == 17
    assert report_of(output)["seed"] == 17
    assert report_of(output)["temperature"] == 0.7


def test_the_cpu_device_offloads_no_layer(tmp_path, model, llm):
    """`device` is also what tells the server a CNE run on the CPU must not
    queue behind a segmentation for the card."""
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert llm.loaded["n_gpu_layers"] == 0


def test_the_cuda_device_offloads_every_layer(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                 device="cuda")

    assert llm.loaded["n_gpu_layers"] == -1


# --------------------------------------------------------------------------
# `device` says what was asked for; `gpu_offload` says what happened
#
# The argument was declared and passed through as `n_gpu_layers=-1` long before
# a build that could honour it was installed. On the CPU wheel that offloads
# nothing at all, in silence -- and the server, which reads `device` to decide
# who takes the card, held one of MAX_CONCURRENT_GPU_JOBS for a run that could
# not use it. Both halves are reported now.
# --------------------------------------------------------------------------

def test_a_cuda_run_on_a_cuda_build_reports_the_offload(tmp_path, model, llm):
    llm.gpu_offload = True
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                          device="cuda")

    report = report_of(output)
    assert report["device"] == "cuda"
    assert report["gpu_offload"] is True
    assert report["warnings"] == []


def test_a_cuda_run_on_a_cpu_build_is_reported_as_ignored(tmp_path, model, llm):
    """The honesty defect: this used to report `device: cuda` over a run done
    entirely on the CPU, with nothing anywhere saying so."""
    llm.gpu_offload = False
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                          device="cuda")

    report = report_of(output)
    assert report["device"] == "cuda"
    assert report["gpu_offload"] is False
    assert any("no GPU support" in warning for warning in report["warnings"])
    assert any("llama-cpp-python" in warning for warning in report["warnings"])


def test_a_cuda_request_a_cpu_build_cannot_honour_is_not_refused(tmp_path, model, llm):
    """Reported, never raised. The server fills `device` in from its own
    setting when the caller sends none, so refusing would fail every request
    reaching a CPU deployment that is configured as a GPU one."""
    llm.gpu_offload = False
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                          device="cuda")

    assert entry_for(output, "a.txt")["status"] == "ok"


def test_a_cpu_run_never_claims_the_gpu_however_capable_the_build(tmp_path, model, llm):
    llm.gpu_offload = True
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                          device="cpu")

    report = report_of(output)
    assert report["gpu_offload"] is False
    assert report["warnings"] == []


def test_the_build_is_asked_about_the_gpu_before_the_weights_are_loaded(
        tmp_path, model, llm, monkeypatch):
    """4.4 GB takes long enough that a deployment which cannot honour `cuda`
    must say so first, not after the model is in memory."""
    order = []
    real_load = sadt_cne.load_model

    monkeypatch.setattr(sadt_cne, "supports_gpu_offload",
                        lambda: order.append("asked") or False)
    monkeypatch.setattr(sadt_cne, "load_model",
                        lambda *a, **k: (order.append("loaded"), real_load(*a, **k))[1])
    write_txt(tmp_path / "notes" / "a.txt")

    sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out",
                 device="cuda")

    assert order == ["asked", "loaded"]


# --------------------------------------------------------------------------
# The CUDA runtime the wheel links and does not ship
#
# `auditwheel` excludes the CUDA runtime from a wheel by policy, so the CUDA
# build of llama-cpp-python needs `libcudart.so.12` and `libcublas.so.12` from
# somewhere else. A workstation with a CUDA toolkit has them and hides the
# problem; the deployment image is `python:3.13-slim` and does not, where
# `import llama_cpp` dies before a line of this tool runs.
# --------------------------------------------------------------------------

def _fake_cuda_wheels(tmp_path, monkeypatch, names):
    """A `nvidia/<package>/lib/<library>` tree, found the way the real one is."""
    root = tmp_path / "site-packages" / "nvidia"
    for package, library in names:
        directory = root / package.split(".")[-1] / "lib"
        directory.mkdir(parents=True, exist_ok=True)
        (directory / library).write_bytes(b"")

    real_find_spec = dependencies.importlib.util.find_spec

    def find_spec(name, *args, **kwargs):
        if name.startswith("nvidia"):
            location = root / name.split(".")[-1]
            if not location.is_dir():
                raise ModuleNotFoundError(name)
            return types.SimpleNamespace(
                submodule_search_locations=[str(location)]
            )
        return real_find_spec(name, *args, **kwargs)

    monkeypatch.setattr(dependencies.importlib.util, "find_spec", find_spec)
    return root


def test_the_cuda_runtime_is_opened_from_the_tools_own_virtualenv(
        tmp_path, monkeypatch):
    _fake_cuda_wheels(tmp_path, monkeypatch, dependencies.CUDA_RUNTIME_LIBRARIES)
    opened = []
    monkeypatch.setattr(ctypes, "CDLL",
                        lambda path, mode=None: opened.append(path))

    returned = dependencies.preload_cuda_runtime()

    assert [Path(path).name for path in returned] == [
        library for _, library in dependencies.CUDA_RUNTIME_LIBRARIES
    ]
    assert opened == returned


def test_the_cuda_runtime_is_opened_globally_so_llama_cpp_resolves_against_it(
        tmp_path, monkeypatch):
    """RTLD_GLOBAL is the whole mechanism: `libggml-cuda.so` is dlopen'd later,
    by llama.cpp, and resolves its undefined symbols against what is already
    loaded globally. Opened privately it would resolve nothing."""
    _fake_cuda_wheels(tmp_path, monkeypatch, dependencies.CUDA_RUNTIME_LIBRARIES)
    modes = []
    monkeypatch.setattr(ctypes, "CDLL",
                        lambda path, mode=None: modes.append(mode))

    dependencies.preload_cuda_runtime()

    assert modes == [ctypes.RTLD_GLOBAL] * len(dependencies.CUDA_RUNTIME_LIBRARIES)


def test_cublaslt_is_opened_before_cublas_which_is_linked_against_it():
    names = [library for _, library in dependencies.CUDA_RUNTIME_LIBRARIES]
    assert names.index("libcublasLt.so.12") < names.index("libcublas.so.12")


def test_a_deployment_without_the_cuda_wheels_preloads_nothing_and_raises_nothing(
        tmp_path, monkeypatch):
    """The CPU build needs none of this, and a workstation whose loader finds
    the system copy needs none of it either. Neither may fail here."""
    _fake_cuda_wheels(tmp_path, monkeypatch, [])

    assert dependencies.preload_cuda_runtime() == []


def test_a_cuda_library_that_will_not_open_is_skipped_not_raised(
        tmp_path, monkeypatch):
    """llama.cpp is about to try the system copy, and the error it raises then
    names the library far more usefully than anything raised here could."""
    _fake_cuda_wheels(tmp_path, monkeypatch, dependencies.CUDA_RUNTIME_LIBRARIES)

    def refuse(path, mode=None):
        raise OSError(f"{Path(path).name}: file too short")

    monkeypatch.setattr(ctypes, "CDLL", refuse)

    assert dependencies.preload_cuda_runtime() == []


def test_the_cuda_runtime_is_preloaded_before_llama_cpp_is_imported(monkeypatch):
    """Order, not presence: the import is what fails without the preload, so
    doing it before the first GPU call would already be too late."""
    order = []
    monkeypatch.setattr(extraction, "preload_cuda_runtime",
                        lambda: order.append("preloaded") or [])
    monkeypatch.setattr(extraction, "require",
                        lambda module, package: order.append("imported"))

    extraction.engine_module()

    assert order == ["preloaded", "imported"]


def test_the_build_reports_whether_it_can_offload(monkeypatch):
    monkeypatch.setattr(extraction, "engine_module",
                        lambda: types.SimpleNamespace(
                            llama_supports_gpu_offload=lambda: 1))
    assert extraction.supports_gpu_offload() is True

    monkeypatch.setattr(extraction, "engine_module",
                        lambda: types.SimpleNamespace(
                            llama_supports_gpu_offload=lambda: 0))
    assert extraction.supports_gpu_offload() is False


def test_the_finish_reason_is_reported_for_every_note(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert entry_for(output, "a.txt")["finish_reason"] == "stop"


def test_an_answer_shaped_wrong_is_a_readable_error():
    class Broken:
        def create_chat_completion(self, **_kwargs):
            return {"choices": []}

    with pytest.raises(ValueError, match="cannot read"):
        extraction.complete(Broken(), [], 10, 0.0)


# --------------------------------------------------------------------------
# Parsing the answer
# --------------------------------------------------------------------------

def test_json_is_found_inside_surrounding_prose():
    answer = 'Sure! Here is the extraction:\n{"pain": "yes"}\nHope that helps.'

    assert extraction.parse_extraction(answer) == {"pain": "yes"}


def test_an_extraction_wrapper_key_is_unwrapped():
    answer = '{"extraction": {"pain": "yes"}}'

    assert extraction.parse_extraction(answer) == {"pain": "yes"}


@pytest.mark.parametrize("answer", [
    "", "No structured findings.", "{not json}", '["pain", "yes"]',
    '{"pain": "yes"', "42", "{}",
])
def test_an_answer_that_is_not_an_object_yields_no_extraction(answer):
    assert extraction.parse_extraction(answer) is None


def test_an_answer_without_json_is_a_failure_not_a_raw_blob(tmp_path, model, llm):
    """Upstream fell back to writing the raw answer, which lands in the output
    folder looking exactly like a successful extraction."""
    write_txt(tmp_path / "notes" / "a.txt")
    write_txt(tmp_path / "notes" / "b.txt")
    llm.answers = ["I am unable to help with that.", ANSWER]

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert not (output / "Extraction_a.txt.txt").exists()
    assert "no JSON object" in entry_for(output, "a.txt")["reason"]


def test_a_list_value_is_rendered_readably():
    rendered = extraction.render_extraction({"symptoms": ["clicking", "pain"]})

    assert rendered == "symptoms : clicking, pain\n"


def test_a_nested_value_stays_json_rather_than_a_python_repr():
    rendered = extraction.render_extraction({"tmj": {"left": 7}})

    assert rendered == 'tmj : {"left": 7}\n'


def test_booleans_and_nulls_render_without_python_spelling():
    rendered = extraction.render_extraction({"pain": True, "surgery": None})

    assert rendered == "pain : true\nsurgery : \n"


def test_the_field_count_is_reported(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert entry_for(output, "a.txt")["fields"] == 2


# --------------------------------------------------------------------------
# The model file
# --------------------------------------------------------------------------

def test_a_bundle_directory_resolves_to_its_single_gguf(tmp_path):
    """`DATA/CNE/models/TMJ/` is a folder, because that is how the manifest
    stages it; the server hands over the folder, not the file inside it."""
    bundle = tmp_path / "TMJ"
    bundle.mkdir()
    weights = bundle / "qwen-ft-q4_k_m.gguf"
    weights.write_bytes(b"x")

    assert extraction.resolve_model_file(bundle) == weights


def test_a_bundle_directory_is_searched_recursively(tmp_path):
    bundle = tmp_path / "TMJ"
    (bundle / "fold_0").mkdir(parents=True)
    weights = bundle / "fold_0" / "model.gguf"
    weights.write_bytes(b"x")

    assert extraction.resolve_model_file(bundle) == weights


def test_a_bundle_holding_several_models_is_refused_rather_than_guessed(tmp_path):
    """Which model vintage ran must never depend on sort order."""
    bundle = tmp_path / "TMJ"
    bundle.mkdir()
    (bundle / "a.gguf").write_bytes(b"x")
    (bundle / "b.gguf").write_bytes(b"x")

    with pytest.raises(ValueError, match="a.gguf, b.gguf"):
        extraction.resolve_model_file(bundle)


def test_a_bundle_holding_no_model_is_a_file_not_found_error(tmp_path):
    bundle = tmp_path / "TMJ"
    bundle.mkdir()
    (bundle / "README.md").write_text("nothing here")

    with pytest.raises(FileNotFoundError, match="setup-models.sh"):
        extraction.resolve_model_file(bundle)


def test_a_gguf_file_is_accepted_directly(tmp_path):
    weights = tmp_path / "model.gguf"
    weights.write_bytes(b"x")

    assert extraction.resolve_model_file(weights) == weights


def test_a_file_that_is_not_a_gguf_is_refused(tmp_path):
    weights = tmp_path / "model.bin"
    weights.write_bytes(b"x")

    with pytest.raises(ValueError, match="GGUF"):
        extraction.resolve_model_file(weights)


def test_a_missing_model_path_is_a_file_not_found_error(tmp_path):
    with pytest.raises(FileNotFoundError, match="setup-models.sh"):
        extraction.resolve_model_file(tmp_path / "nowhere")


@pytest.mark.parametrize("name,hint", [
    ("TMJ/qwen-ft-q4_k_m.gguf", "TMJ"),
    ("Ortho/model-q4_0.gguf", "Ortho"),
    ("models/generic.gguf", ""),
])
def test_the_model_path_hints_at_the_note_type_it_was_trained_for(name, hint):
    assert extraction.model_type_hint(name) == hint


def test_a_model_named_for_the_other_note_type_is_warned_about_not_refused(tmp_path, model, llm):
    """A file name is not a contract, so this cannot refuse -- but pairing the
    Ortho weights with the TMJ instruction gives an answer that looks right."""
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "Ortho", model, tmp_path / "out")

    warnings = report_of(output)["warnings"]
    assert len(warnings) == 1 and "TMJ" in warnings[0]
    assert report_of(output)["summary"] == "1/1 note(s) extracted"


def test_a_matching_model_produces_no_warning(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert report_of(output)["warnings"] == []


def test_the_model_file_name_is_recorded_in_the_report(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert report_of(output)["model"] == "qwen-ft-q4_k_m.gguf"


# --------------------------------------------------------------------------
# The tally (defect: "All files processed successfully!" printed unconditionally)
# --------------------------------------------------------------------------

def test_the_report_counts_what_was_written_not_what_was_walked_past(tmp_path, model, llm):
    """Upstream computed `successfully_processed`/`failed_files` and then
    printed "All files processed successfully!" regardless (CNE_CLI.py:240-257).
    0 of 40 succeeding still exited 0."""
    notes = tmp_path / "notes"
    write_txt(notes / "a.txt")
    write_txt(notes / "b.txt")
    write_pdf(notes / "c.pdf", text="")

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert report_of(output)["summary"] == "2/3 note(s) extracted"


def test_a_run_where_nothing_succeeded_raises_instead_of_exiting_zero(tmp_path, model, llm):
    write_pdf(tmp_path / "notes" / "a.pdf", text="")

    with pytest.raises(ValueError, match="extracted nothing"):
        sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")


def test_the_failure_message_names_every_note_and_its_reason(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_pdf(notes / "a.pdf", text="")
    write_pdf(notes / "b.pdf", text="")

    with pytest.raises(ValueError) as raised:
        sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert "a.pdf" in str(raised.value) and "b.pdf" in str(raised.value)


def test_every_note_gets_a_report_entry_whatever_happened(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_txt(notes / "a.txt")
    write_pdf(notes / "b.pdf", text="")

    output = sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert [entry["input"] for entry in report_of(output)["notes"]] == [
        "a.txt", "b.pdf",
    ]


def test_the_report_records_how_long_the_run_took(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert report_of(output)["duration_seconds"] >= 0


def test_run_returns_the_output_directory(tmp_path, model, llm):
    write_txt(tmp_path / "notes" / "a.txt")

    output = sadt_cne.run(tmp_path / "notes", "TMJ", model, tmp_path / "out")

    assert output == tmp_path / "out"


def test_nothing_is_written_outside_the_output_directory(tmp_path, model, llm):
    notes = tmp_path / "notes"
    write_txt(notes / "a.txt")
    before = sorted(path.name for path in notes.iterdir())

    sadt_cne.run(notes, "TMJ", model, tmp_path / "out")

    assert sorted(path.name for path in notes.iterdir()) == before


# --------------------------------------------------------------------------
# stderr (defect: raw fd surgery with no try/finally, CNE/CNE.py:155-169)
# --------------------------------------------------------------------------

def test_quiet_stderr_gives_stderr_back_on_success():
    original = sys.stderr

    with extraction.quiet_stderr():
        pass

    assert sys.stderr is original


def test_quiet_stderr_gives_stderr_back_when_the_load_raises():
    """Upstream's restore sat AFTER the load, so a load that raised left the
    interpreter writing its traceback into /dev/null for the rest of the run."""
    original = sys.stderr

    with pytest.raises(RuntimeError):
        with extraction.quiet_stderr():
            raise RuntimeError("could not load the model")

    assert sys.stderr is original


def test_quiet_stderr_captures_what_the_loader_printed():
    with extraction.quiet_stderr() as captured:
        print("llama_model_loader: loaded meta data", file=sys.stderr)

    assert "llama_model_loader" in captured.getvalue()


# --------------------------------------------------------------------------
# The contract with the server
# --------------------------------------------------------------------------

def test_run_annotations_are_stdlib_types_only():
    """The server publishes the schema by importing this package and reading
    the signature; anything it cannot map is a hard refusal at startup."""
    import typing

    allowed = {Path, str, int, float, bool}
    for name, annotation in typing.get_type_hints(sadt_cne.run).items():
        if typing.get_origin(annotation) is typing.Literal:
            assert all(isinstance(option, str)
                       for option in typing.get_args(annotation)), name
            continue
        assert annotation in allowed, f"{name}: {annotation!r}"


def test_every_argument_is_documented_for_the_panel():
    import inspect

    documented = inspect.getdoc(sadt_cne.run)
    for name in inspect.signature(sadt_cne.run).parameters:
        assert f"{name}:" in documented, name


def test_the_model_argument_is_a_path_named_model():
    """That name and that type are what make the server publish it as a hosted
    NAME the client picks -- never a file a clinician can upload."""
    import inspect

    parameter = inspect.signature(sadt_cne.run).parameters["model"]
    assert parameter.annotation is Path
    assert parameter.default is inspect.Parameter.empty


def test_importing_the_package_loads_no_inference_engine():
    """CI imports it on every pull request to publish the schema, and the
    server regenerates the schema at startup. Neither may pay for llama.cpp."""
    source = (
        "import sys; import sadt_cne; "
        "print([m for m in ('llama_cpp', 'fitz', 'docx') if m in sys.modules])"
    )
    src = str(Path(__file__).resolve().parent.parent / "src")
    result = subprocess.run(
        [sys.executable, "-c", source],
        capture_output=True, text=True, cwd=src, check=True,
    )

    assert result.stdout.strip() == "[]"


def test_the_tool_is_declared_in_its_pyproject():
    project = Path(__file__).resolve().parent.parent / "pyproject.toml"
    text = project.read_text()

    assert "[tool.sadt]" in text
    assert 'name = "CNE"' in text


# --------------------------------------------------------------------------
# Needs a CUDA build and a card. Deselected by default; run it by hand on a
# GPU deployment with `-m gpu -o addopts=`.
# --------------------------------------------------------------------------

@pytest.mark.gpu
def test_the_installed_build_has_gpu_support_compiled_in():
    """Cheap, and it is the whole precondition: without this every `cuda`
    request is answered on the CPU and reported as ignored."""
    assert extraction.supports_gpu_offload(), (
        "the llama-cpp-python installed here is the CPU build. Check the "
        "`[[tool.uv.index]]` url in pyproject.toml."
    )


def _vram_used_by_this_process():
    """MiB this process holds on the card, from nvidia-smi. 0 if it holds none."""
    query = subprocess.run(
        ["nvidia-smi", "--query-compute-apps=pid,used_memory",
         "--format=csv,noheader,nounits"],
        capture_output=True, text=True, check=True,
    )
    for line in query.stdout.splitlines():
        pid, _, used = line.partition(",")
        if pid.strip() == str(os.getpid()):
            return int(used.strip())
    return 0


@pytest.mark.gpu
@pytest.mark.models
def test_the_weights_really_land_on_the_card(tmp_path):
    """That the build CAN offload is not that this run DID. Measured against
    the card rather than inferred from the argument:

        CNE_MODEL=../../DATA/CNE/models/TMJ \
        .venv/bin/python -m pytest -m "gpu and models" -o addopts=
    """
    bundle = os.environ.get("CNE_MODEL")
    if not bundle:
        pytest.skip("set CNE_MODEL to a bundle under DATA/CNE/models/")

    before = _vram_used_by_this_process()
    engine = extraction.load_model(
        extraction.resolve_model_file(Path(bundle)), 2048, 0, -1
    )
    after = _vram_used_by_this_process()
    del engine

    # A 7B model at q4_k_m is ~4.4 GB of weights. Anything above a gigabyte is
    # unambiguous; the threshold is not tuned to a particular quantisation.
    assert after - before > 1024, (
        f"loading with n_gpu_layers=-1 took {after - before} MiB of VRAM, so "
        f"the layers stayed on the CPU"
    )


@pytest.mark.gpu
@pytest.mark.models
def test_a_real_cuda_run_reports_the_offload_and_warns_about_nothing(tmp_path):
    bundle = os.environ.get("CNE_MODEL")
    if not bundle:
        pytest.skip("set CNE_MODEL to a bundle under DATA/CNE/models/")

    write_txt(
        tmp_path / "notes" / "note.txt",
        "Patient reports left TMJ pain, 7/10, worse on chewing. "
        "Clicking on opening. No locking.",
    )

    output = sadt_cne.run(tmp_path / "notes", "TMJ", Path(bundle),
                          tmp_path / "out", device="cuda")

    report = report_of(output)
    assert report["device"] == "cuda"
    assert report["gpu_offload"] is True
    assert report["warnings"] == []


# --------------------------------------------------------------------------
# Needs the real weights. Skipped by default; run it by hand.
# --------------------------------------------------------------------------

@pytest.mark.models
def test_the_real_model_answers_with_json(tmp_path):
    """The only test that loads a 4.4 GB GGUF. Point CNE_MODEL at a bundle:

        CNE_MODEL=../../DATA/CNE/models/TMJ .venv/bin/python -m pytest -m models
    """
    bundle = os.environ.get("CNE_MODEL")
    if not bundle:
        pytest.skip("set CNE_MODEL to a bundle under DATA/CNE/models/")

    write_txt(
        tmp_path / "notes" / "note.txt",
        "Patient reports left TMJ pain, 7/10, worse on chewing. "
        "Clicking on opening. No locking.",
    )

    output = sadt_cne.run(tmp_path / "notes", "TMJ", Path(bundle), tmp_path / "out")

    data = json.loads((output / "Extraction_note.txt.json").read_text())
    assert isinstance(data, dict) and data


# ---------------------------------------------------------------------------
# What the published fine-tune actually answers
#
# Found by running the real 4.4 GB TMJ model through the API: it emits 46
# `key: value` lines and no JSON at all. The first port refused every one of
# them, because upstream reached that shape only through its raw-answer
# fallback -- so what looked like a defect ("a truncated blob written as if it
# were the result") is, for this model, the normal path.
# ---------------------------------------------------------------------------

REAL_TMJ_ANSWER = """patient_id: unknown
patient_age: 34
maximum_opening: 38mm
jaw_locking: false
onset_triggers: motor vehicle accident 3 years ago with whiplash
muscle_pain_location: right masseter | right temporalis
disc_displacement: right tmj anterior displacement with reduction
average_daily_pain_intensity: 6/10
pain_relieving_factors: unknown
pain_relieving_factors: nsaids
"""


def test_the_real_models_key_value_answer_is_an_extraction():
    data = extraction.parse_extraction(REAL_TMJ_ANSWER)
    assert data is not None
    assert data["patient_age"] == "34"
    assert data["maximum_opening"] == "38mm"
    assert data["muscle_pain_location"] == "right masseter | right temporalis"


def test_a_repeated_field_keeps_the_first_value_not_the_last():
    """The published fine-tune repeats `pain_relieving_factors` in one answer.
    A plain dict build keeps whichever came last, silently."""
    data = extraction.parse_extraction(REAL_TMJ_ANSWER)
    assert data["pain_relieving_factors"] == "unknown"


def test_every_field_of_a_long_real_answer_survives():
    data = extraction.parse_extraction(REAL_TMJ_ANSWER)
    assert len(data) == 9          # ten lines, one of them a repeat


def test_json_is_still_preferred_when_the_model_produces_it():
    data = extraction.parse_extraction('Here you go: {"patient_age": 34} hope that helps')
    assert data == {"patient_age": 34}


def test_prose_is_not_mistaken_for_an_extraction():
    """A clinical note is full of `WORD: text` lines. Reading one as fields
    would write the note back out as if the model had extracted it."""
    note = (
        "CHIEF COMPLAINT: Right TMJ pain for 8 months.\n"
        "HISTORY: 34-year-old patient reports clicking on opening.\n"
        "EXAM: Maximum unassisted opening 38 mm. Deviation to the right.\n"
        "IMAGING: MRI shows anterior disc displacement with reduction.\n"
    )
    assert extraction.parse_extraction(note) is None


def test_an_uppercase_heading_is_not_a_field():
    assert extraction.parse_extraction("ASSESSMENT: right TMJ derangement\n") is None


def test_one_field_line_buried_in_prose_is_not_an_extraction():
    answer = (
        "I looked at the note and here is what I found.\n"
        "The patient is thirty-four years old and reports pain.\n"
        "patient_age: 34\n"
        "That is all I could determine from the text provided.\n"
    )
    assert extraction.parse_extraction(answer) is None


def test_a_field_name_written_with_spaces_becomes_an_identifier():
    data = extraction.parse_extraction("patient age: 34\njaw locking: false\n")
    assert data == {"patient_age": "34", "jaw_locking": "false"}


def test_a_value_containing_a_colon_keeps_all_of_it():
    data = extraction.parse_extraction(
        "pain_onset_date: 2019-04-01\nnote: seen at 09:30 by Dr. A\ncomment: ok\n")
    assert data["note"] == "seen at 09:30 by Dr. A"


def test_a_field_with_an_empty_value_is_not_a_field():
    assert extraction.parse_extraction("patient_age:\njaw_locking:\n") is None


def test_an_empty_answer_is_no_extraction():
    assert extraction.parse_extraction("") is None
    assert extraction.parse_extraction("   \n\n  ") is None


def test_a_single_object_in_an_array_is_read_as_that_object():
    """Upstream slices between the first `{` and the last `}` because the
    fine-tunes wrap their JSON in a sentence. A one-element array falls out of
    that slicing, and reading it is right."""
    assert extraction.parse_extraction('[{"patient_age": 34}]') == {"patient_age": 34}


def test_an_array_of_several_objects_is_refused_rather_than_spanned():
    """The same slicing over two objects gives `{...}, {...}`, which is not
    valid JSON -- so it is refused rather than half-read."""
    assert extraction.parse_extraction('[{"patient_age": 34}, {"patient_age": 51}]') is None


def test_blank_lines_between_fields_do_not_break_the_ratio():
    data = extraction.parse_extraction("patient_age: 34\n\n\njaw_locking: false\n")
    assert data == {"patient_age": "34", "jaw_locking": "false"}
